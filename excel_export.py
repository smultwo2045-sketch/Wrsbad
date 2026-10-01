"""Excel export: shifts grouped by manager."""
from __future__ import annotations

import re
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from address import normalize_address, shorten_address
from helpers import fmt_hours

MSK = timezone(timedelta(hours=3))


def msk_now() -> datetime:
    return datetime.now(MSK)


def _safe_sheet_name(name: str, used: set) -> str:
    s = re.sub(r'[:\\/?*\[\]]', "-", str(name or "Менеджер"))
    s = s.strip()[:31] or "Менеджер"
    base, n = s, 1
    while s in used:
        suffix = f"_{n}"
        s = (base[: 31 - len(suffix)] + suffix)
        n += 1
    used.add(s)
    return s


def _fmt_time(val) -> str:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ""
    if isinstance(val, datetime):
        return val.strftime("%H:%M")
    if hasattr(val, "strftime"):
        try:
            return val.strftime("%H:%M")
        except Exception:
            pass
    s = str(val).strip()
    if not s or s.lower() == "nan":
        return ""
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y %H:%M",
        "%Y-%m-%d %H:%M",
        "%H:%M:%S",
        "%H:%M",
    ):
        try:
            return datetime.strptime(s[:19], fmt).strftime("%H:%M")
        except Exception:
            continue
    # pandas Timestamp string
    if " " in s and ":" in s:
        try:
            return s.split(" ")[1][:5]
        except Exception:
            pass
    return s[:5] if len(s) >= 5 else s


def build_shifts_excel(
    shift_details: pd.DataFrame,
    shift_adj: pd.DataFrame,
    ab_df: pd.DataFrame,
    shift_ctx: dict,
    manager_order: Optional[List[str]] = None,
    regions: Optional[dict] = None,
) -> bytes:
    regions = regions or {}
    sc = (shift_ctx or {}).get("sc")
    ec = (shift_ctx or {}).get("ec")
    report_date = (shift_ctx or {}).get("report_date") or msk_now().strftime("%d.%m.%Y")
    title_date = f"Смены за {report_date}"

    if not isinstance(shift_details, pd.DataFrame) or shift_details.empty:
        raise ValueError("Нет данных смен для экспорта")

    ab = ab_df.copy() if isinstance(ab_df, pd.DataFrame) else pd.DataFrame(columns=["Адрес магазина", "Менеджер"])
    if ab.empty:
        ab = pd.DataFrame(columns=["Адрес магазина", "Менеджер"])
    if "Адрес магазина" not in ab.columns:
        ab["Адрес магазина"] = ""
    if "Менеджер" not in ab.columns:
        ab["Менеджер"] = ""

    ab["_n"] = ab["Адрес магазина"].fillna("").astype(str).apply(normalize_address)
    ab["_short"] = ab["Адрес магазина"].fillna("").astype(str).apply(shorten_address)
    mgr_map = (
        ab.drop_duplicates("_n").set_index("_n")["Менеджер"].to_dict()
        if not ab.empty
        else {}
    )

    inner = shift_details.copy()
    if "Адрес магазина" not in inner.columns:
        # try alternatives
        for c in ("Адрес объекта", "Адрес"):
            if c in inner.columns:
                inner["Адрес магазина"] = inner[c]
                break
        else:
            inner["Адрес магазина"] = ""
    if "_n" not in inner.columns:
        inner["_n"] = inner["Адрес магазина"].fillna("").astype(str).apply(normalize_address)
    inner["Менеджер"] = inner["_n"].map(mgr_map).fillna("Не найден")
    inner["_short"] = inner["Адрес магазина"].fillna("").astype(str).apply(
        lambda x: shorten_address(str(x))
    )
    if "_h" not in inner.columns:
        inner["_h"] = None

    adj = shift_adj.copy() if isinstance(shift_adj, pd.DataFrame) else pd.DataFrame()
    if not adj.empty:
        if "Адрес магазина" not in adj.columns:
            for c in ("Адрес объекта", "Адрес"):
                if c in adj.columns:
                    adj["Адрес магазина"] = adj[c]
                    break
        if "_n" not in adj.columns and "Адрес магазина" in adj.columns:
            adj["_n"] = adj["Адрес магазина"].fillna("").astype(str).apply(normalize_address)
        if "Менеджер" not in adj.columns and "_n" in adj.columns:
            adj["Менеджер"] = adj["_n"].map(mgr_map).fillna("Не найден")
        if "_short" not in adj.columns and "Адрес магазина" in adj.columns:
            adj["_short"] = adj["Адрес магазина"].fillna("").astype(str).apply(
                lambda x: shorten_address(str(x))
            )

    nonauth_by_mgr: Dict[str, List[str]] = {}
    adj_by_mgr: Dict[str, List[str]] = {}
    if not ab.empty and "Менеджер" in ab.columns:
        for mgr, ab_grp in ab.groupby("Менеджер"):
            if not str(mgr).strip():
                continue
            ab_set = set(ab_grp["_n"].unique())
            inner_set = set(inner[inner["Менеджер"] == mgr]["_n"].unique())
            missing = ab_set - inner_set
            if missing:
                nonauth_by_mgr[str(mgr)] = ab_grp[ab_grp["_n"].isin(missing)]["_short"].tolist()
            if not adj.empty and "Менеджер" in adj.columns:
                aj = adj[adj["Менеджер"] == mgr]
                if not aj.empty:
                    col = "_short" if "_short" in aj.columns else "Адрес магазина"
                    adj_by_mgr[str(mgr)] = [
                        shorten_address(str(x)) for x in aj[col].tolist()
                    ]

    order = [str(m) for m in (manager_order or [])]
    mgr_list: List[str] = []
    for m in order:
        if (
            m in inner["Менеджер"].astype(str).values
            or m in adj_by_mgr
            or m in nonauth_by_mgr
        ):
            mgr_list.append(m)
    for m in list(inner["Менеджер"].astype(str).unique()) + list(nonauth_by_mgr) + list(adj_by_mgr):
        if m and m not in mgr_list:
            mgr_list.append(m)
    if not mgr_list:
        mgr_list = ["Без менеджера"]

    def region_of(m: str) -> str:
        for reg, ms in regions.items():
            if m in ms:
                return str(reg)
        return ""

    NAVY, WHITE, ALT = "0A246A", "FFFFFF", "F2F4FB"
    RED_L, RED_H, GREEN_L, GREEN_H, BC = "FFF0F0", "C0392B", "F0FFF4", "1A6B3C", "B8BCC8"
    AMBER = "FFF8E1"

    def bd(c=BC):
        s = Side(style="thin", color=c)
        return Border(left=s, right=s, top=s, bottom=s)

    def hdr(ws, r, c, val, color=NAVY, align="center"):
        cell = ws.cell(row=r, column=c, value=val)
        cell.font = Font(name="Calibri", bold=True, color=WHITE, size=9)
        cell.fill = PatternFill("solid", start_color=color)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        cell.border = bd(color)
        return cell

    def dat(ws, r, c, val, bg=WHITE, align="center", wrap=False, bold=False, col_txt="000000"):
        cell = ws.cell(row=r, column=c, value=val if val is not None else "")
        cell.font = Font(name="Calibri", size=9, bold=bold, color=col_txt)
        cell.fill = PatternFill("solid", start_color=bg)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        cell.border = bd()
        return cell

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    # Summary
    ws = wb.create_sheet("Сводка")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    cell = ws["A1"]
    cell.value = title_date
    cell.font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.fill = PatternFill("solid", start_color="EEF2FA")
    ws.row_dimensions[1].height = 28

    headers = [
        "Менеджер",
        "Регион",
        "В адреске",
        "Открыли (магазин)",
        "Не авт.",
        "Прилегающая",
        "Часы (магазин)",
        "Коротких <3ч",
    ]
    for i, h in enumerate(headers, 1):
        hdr(ws, 3, i, h, align="center" if i > 1 else "left")

    ri = 4
    for mi, mgr in enumerate(mgr_list):
        bg = ALT if mi % 2 else WHITE
        sub = inner[inner["Менеджер"].astype(str) == mgr]
        ab_n = int(ab[ab["Менеджер"].astype(str) == mgr]["_n"].nunique()) if not ab.empty else 0
        opened = int(sub["_n"].nunique()) if not sub.empty else 0
        na_n = len(nonauth_by_mgr.get(mgr, []))
        aj_n = len(adj_by_mgr.get(mgr, []))
        try:
            hours = float(pd.to_numeric(sub["_h"], errors="coerce").fillna(0).sum()) if not sub.empty else 0.0
            short_n = int((pd.to_numeric(sub["_h"], errors="coerce") < 3).sum()) if not sub.empty else 0
        except Exception:
            hours, short_n = 0.0, 0
        vals = [mgr, region_of(mgr), ab_n, opened, na_n, aj_n, fmt_hours(hours), short_n]
        for ci, v in enumerate(vals, 1):
            col_txt = "000000"
            if ci == 5 and na_n:
                col_txt = RED_H
            if ci == 6 and aj_n:
                col_txt = GREEN_H
            dat(ws, ri, ci, v, bg=bg, align="left" if ci == 1 else "center", bold=(ci == 1), col_txt=col_txt)
        ri += 1

    for i, w in enumerate([28, 12, 12, 14, 10, 12, 14, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    used_names = {"Сводка"}
    for mgr in mgr_list:
        sub_i = inner[inner["Менеджер"].astype(str) == mgr].copy()
        if not sub_i.empty and "_h" in sub_i.columns:
            try:
                sub_i = sub_i.assign(_hsort=pd.to_numeric(sub_i["_h"], errors="coerce")).sort_values(
                    "_hsort", ascending=True
                )
            except Exception:
                pass
        na_list = nonauth_by_mgr.get(mgr, [])
        adj_list = adj_by_mgr.get(mgr, [])
        region = region_of(mgr)
        ab_total = int(ab[ab["Менеджер"].astype(str) == mgr]["_n"].nunique()) if not ab.empty else 0
        shift_n = int(sub_i["_n"].nunique()) if not sub_i.empty else 0
        try:
            total_h = float(pd.to_numeric(sub_i["_h"], errors="coerce").fillna(0).sum()) if not sub_i.empty else 0.0
            short_n = int((pd.to_numeric(sub_i["_h"], errors="coerce") < 3).sum()) if not sub_i.empty else 0
        except Exception:
            total_h, short_n = 0.0, 0

        ws = wb.create_sheet(title=_safe_sheet_name(mgr, used_names))
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:D1")
        c = ws["A1"]
        c.value = f"{mgr}  ({region})  —  {title_date}" if region else f"{mgr}  —  {title_date}"
        c.font = Font(name="Calibri", bold=True, size=12, color=NAVY)
        c.fill = PatternFill("solid", start_color="EEF2FA")
        ws.row_dimensions[1].height = 26

        ws.merge_cells("A2:D2")
        c2 = ws["A2"]
        na_part = f"   |   Не авт.: {len(na_list)}" if na_list else ""
        adj_part = f"   |   Прилег.: {len(adj_list)}" if adj_list else ""
        c2.value = (
            f"В адреске: {ab_total}   |   Открыли: {shift_n}{na_part}{adj_part}"
            f"   |   Итого: {fmt_hours(total_h)}   |   Коротких: {short_n}"
        )
        c2.font = Font(name="Calibri", size=8, color="444444")
        c2.fill = PatternFill("solid", start_color="EEF2FA")

        G1, NC_COL, G2, AJ = 5, 6, 7, 8
        hdr(ws, 4, 1, "Адрес магазина", align="left")
        hdr(ws, 4, 2, "Начало")
        hdr(ws, 4, 3, "Конец")
        hdr(ws, 4, 4, "Длительность")
        if na_list:
            hdr(ws, 4, NC_COL, f"Не авторизованы — {len(na_list)}", RED_H, "left")
        if adj_list:
            hdr(ws, 4, AJ, f"Прилегающая — {len(adj_list)}", GREEN_H, "left")

        max_rows = max(len(sub_i), len(na_list), len(adj_list), 1) + 10
        for rr in range(1, max_rows + 1):
            ws.cell(row=rr, column=G1).fill = PatternFill("solid", start_color="F4F4F4")
            if adj_list:
                ws.cell(row=rr, column=G2).fill = PatternFill("solid", start_color="F4F4F4")

        ri = 5
        for _, row in sub_i.iterrows():
            try:
                h_val = row.get("_h")
                h_num = float(h_val) if h_val is not None and not pd.isna(h_val) else None
            except Exception:
                h_num = None
            is_short = h_num is not None and h_num < 3
            bg = AMBER if is_short else (ALT if (ri % 2 == 0) else WHITE)
            addr = row.get("_short") or shorten_address(str(row.get("Адрес магазина", "")))
            start_v = _fmt_time(row.get(sc)) if sc and sc in row.index else ""
            end_v = _fmt_time(row.get(ec)) if ec and ec in row.index else ""
            dur = fmt_hours(h_num) if h_num is not None else ""
            dat(ws, ri, 1, addr, bg=bg, align="left")
            dat(ws, ri, 2, start_v, bg=bg)
            dat(ws, ri, 3, end_v, bg=bg)
            dat(ws, ri, 4, dur, bg=bg, bold=is_short, col_txt=RED_H if is_short else "000000")
            ri += 1

        for i, addr in enumerate(na_list):
            dat(ws, 5 + i, NC_COL, addr, bg=RED_L, align="left", col_txt=RED_H)
        for i, addr in enumerate(adj_list):
            dat(ws, 5 + i, AJ, addr, bg=GREEN_L, align="left", col_txt=GREEN_H)

        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 10
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions[get_column_letter(G1)].width = 2
        ws.column_dimensions[get_column_letter(NC_COL)].width = 32
        ws.column_dimensions[get_column_letter(G2)].width = 2
        ws.column_dimensions[get_column_letter(AJ)].width = 32

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
